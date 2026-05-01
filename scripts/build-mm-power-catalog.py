from __future__ import annotations

import argparse
import json
import re
import unicodedata
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


TRUTHY = {"sim", "yes", "true", "1"}
EMPTY = {"", "-", "—", "nao", "não", "none", "null", "nan"}


def clean(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return "" if norm_text(text) in EMPTY else text


def norm_text(value) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return text.strip().lower()


def slugify(value) -> str:
    text = norm_text(value)
    text = re.sub(r"[^a-z0-9]+", "-", text).strip("-")
    return text or "unknown"


def is_yes(value) -> bool:
    return norm_text(value) in TRUTHY


def split_top_level(text: str, sep: str = ";") -> list[str]:
    parts: list[str] = []
    cur: list[str] = []
    round_depth = 0
    square_depth = 0
    for ch in text:
        if ch == "(":
            round_depth += 1
        elif ch == ")" and round_depth:
            round_depth -= 1
        elif ch == "[":
            square_depth += 1
        elif ch == "]" and square_depth:
            square_depth -= 1

        if ch == sep and round_depth == 0 and square_depth == 0:
            part = "".join(cur).strip()
            if part:
                parts.append(part)
            cur = []
        else:
            cur.append(ch)
    part = "".join(cur).strip()
    if part:
        parts.append(part)
    return parts


def merge_modifier_only_segments(segments: list[str]) -> list[str]:
    merged: list[str] = []
    for segment in segments:
        raw = segment.strip()
        is_bracket_note = re.match(r"^\[?(?:Extra|Flaw|Custom)\s*:", raw, re.I)
        is_extra_note = re.match(r"^(?:Precise|Subtle|Accurate\b|Improved Critical\b|Multiattack\b)", raw, re.I)
        is_flaw_note = re.match(r"^(?:Side Effect\b|Quirk\b|Recharge\b|Cooldown\b|Charge\b)", raw, re.I)
        if merged and (is_bracket_note or is_extra_note or is_flaw_note):
            if raw.startswith("[") and raw.endswith("]"):
                merged[-1] = f"{merged[-1]} {raw}"
            elif is_flaw_note:
                merged[-1] = f"{merged[-1]} [Flaw: {raw}]"
            elif is_extra_note:
                merged[-1] = f"{merged[-1]} [Extra: {raw}]"
            else:
                merged[-1] = f"{merged[-1]} [{raw}]"
        else:
            merged.append(raw)
    return merged


def normalize_resistance(value: str) -> str:
    raw = norm_text(value)
    if not raw:
        return ""
    if raw in {"thg", "tgh", "toughness", "resistencia", "resistance"}:
        return "thg"
    if raw in {"fort", "fortitude"}:
        return "fort"
    if raw in {"will", "vontade"}:
        return "will"
    if raw in {"dodge", "esquiva"}:
        return "dodge"
    if raw in {"parry", "aparar"}:
        return "parry"
    if raw in {"stgr", "str", "strength", "forca"}:
        return "stgr"
    return raw.replace(" ", "_")


def normalize_stat(value: str) -> str:
    raw = norm_text(value).replace(".", "")
    aliases = {
        "attack": "stgr",
        "atk": "stgr",
        "strength": "stgr",
        "stgr": "stgr",
        "sp atk": "int",
        "spatk": "int",
        "special attack": "int",
        "intellect": "int",
        "intelect": "int",
        "int": "int",
        "defense": "thg",
        "def": "thg",
        "sp def": "will",
        "spdef": "will",
        "special defense": "will",
        "speed": "speed",
        "initiative": "initiative",
        "evasion": "dodge",
        "accuracy": "acerto",
        "pre": "presence",
    }
    return aliases.get(raw, normalize_resistance(raw) or raw.replace(" ", "_"))


def parse_rank(segment: str, effect_name: str, default_source: str = "move") -> dict:
    prefix = re.escape(effect_name)
    match = re.search(rf"\b{prefix}\s+(-?\d+)\b", segment, re.I)
    if match:
        return {"source": "fixed", "value": int(match.group(1))}
    match = re.search(r"\bRank\s*=\s*(-?\d+)\b", segment, re.I)
    if match:
        return {"source": "fixed", "value": int(match.group(1))}
    return {"source": default_source, "value": None}


def parse_range(row: dict, segment: str) -> str:
    text = norm_text(segment)
    if "perception" in text:
        return "perception"
    if "[ranged]" in text or is_yes(row.get("Ranged")):
        return "ranged"
    if "[close]" in text:
        return "close"
    return "close"


def parse_area(row: dict, segment: str) -> dict | None:
    raw = clean(row.get("Área"))
    text = norm_text(segment)
    area_type = ""
    if raw:
        area_type = raw
    match = re.search(r"\[Area:\s*([^\]]+)\]", segment, re.I)
    if match:
        area_type = match.group(1).strip()
    perception = is_yes(row.get("Perception Area")) or "perception area" in text
    if not area_type and not perception:
        return None
    return {
        "type": area_type or "Perception",
        "perception": perception,
        "selective": "selective" in text,
    }


def parse_modifiers(segment: str) -> tuple[list[str], list[str], dict]:
    extras: list[str] = []
    flaws: list[str] = []
    flags = {
        "linked": bool(re.search(r"\bLinked\b", segment, re.I)),
        "secondaryEffect": bool(re.search(r"\bSecondary Effect\b", segment, re.I)),
        "fades": bool(re.search(r"\bFades\b", segment, re.I)),
        "unreliable": bool(re.search(r"\bUnreliable\b", segment, re.I)),
        "reaction": bool(re.search(r"\bReaction\b", segment, re.I)),
        "selective": bool(re.search(r"\bSelective\b", segment, re.I)),
        "limited": bool(re.search(r"\bLimited\b", segment, re.I)),
        "multiattack": bool(re.search(r"\bMultiattack\b", segment, re.I)),
    }
    for match in re.finditer(r"\[Extra:\s*([^\]]+)\]", segment, re.I):
        extras.append(match.group(1).strip())
    for match in re.finditer(r"\[Flaw:\s*([^\]]+)\]", segment, re.I):
        flaws.append(match.group(1).strip())
    for match in re.finditer(r"\[(?!(?:Extra|Flaw)\s*:)(?:Custom:\s*)?([^]\[]*?\b(?:Affects Others|Precise|Subtle|Indirect|Triggered|Area|Limited|Increased Range|Increased Duration|Secondary Effect|Variable Descriptor|Alternate Resistance)[^]\[]*)\]", segment, re.I):
        label = match.group(1).strip()
        if not label:
            continue
        if re.search(r"\bLimited\b", label, re.I):
            if label not in flaws:
                flaws.append(label)
        elif label not in extras:
            extras.append(label)
    for key, label in {
        "multiattack": "Multiattack",
        "secondaryEffect": "Secondary Effect",
        "reaction": "Reaction",
        "selective": "Selective",
    }.items():
        if flags[key] and label not in extras:
            extras.append(label)
    if flags["unreliable"] and "Unreliable" not in flaws:
        flaws.append("Unreliable")
    if flags["fades"] and "Fades" not in flaws:
        flaws.append("Fades")
    if flags["limited"] and not any(norm_text(x).startswith("limited") for x in flaws):
        flaws.append("Limited")
    return extras, flaws, flags


def parse_resistance(row: dict, segment: str, fallback: str = "") -> str:
    alt = re.search(r"Alternate Resistance:\s*([A-Za-z ]+)", segment, re.I)
    if alt:
        return normalize_resistance(alt.group(1))
    resisted = re.search(r"Resisted by\s+([A-Za-z ]+)", segment, re.I)
    if resisted:
        return normalize_resistance(resisted.group(1))
    return normalize_resistance(clean(row.get("Resist Stat")) or fallback)


def parse_affliction_conditions(segment: str, status_hint: str) -> list[dict]:
    status = norm_text(status_hint)
    if re.match(r"^(?:Linked\s+)?Dazzle\b", segment, re.I):
        return [
            {"degree": 1, "condition": "impaired"},
            {"degree": 2, "condition": "disabled"},
            {"degree": 3, "condition": "unaware"},
        ]
    status_defaults = {
        "paralyze": ["fatigued", "immobile", "paralyzed"],
        "freeze": ["dazed", "stunned", "transformed"],
        "confusion": ["vulnerable", "defenseless", "controlled"],
        "sleep": ["fatigued", "exhausted", "asleep"],
        "poison": ["impaired", "disabled", "incapacitated"],
        "burn": ["impaired", "disabled", "incapacitated"],
        "flinch": ["dazed", "stunned", "incapacitated"],
        "trapped": ["hindered", "immobile", "paralyzed"],
    }
    if status in status_defaults:
        return [{"degree": i + 1, "condition": cond} for i, cond in enumerate(status_defaults[status])]

    inside = ""
    match = re.search(r"Affliction\s+\d*\s*\(([^)]*)\)", segment, re.I)
    if match:
        inside = match.group(1)
    if inside:
        first_clause = re.split(r";|Resisted by|Custom:", inside, flags=re.I)[0]
        tokens = [re.sub(r"\[[^\]]+\]", "", x).strip() for x in first_clause.split(",")]
        tokens = [slugify(x) for x in tokens if x.strip()]
        if tokens:
            while len(tokens) < 3:
                tokens.append(tokens[-1])
            return [{"degree": i + 1, "condition": tokens[i]} for i in range(3)]

    return [
        {"degree": 1, "condition": "dazed"},
        {"degree": 2, "condition": "stunned"},
        {"degree": 3, "condition": "incapacitated"},
    ]


def effect_type_from_segment(segment: str) -> tuple[str, str]:
    text = segment.strip()
    text = re.sub(r"^Linked\s+", "", text, flags=re.I).strip()
    text = re.sub(r"^AE\d+\s*(?:\([^)]*\))?\s*:\s*", "", text, flags=re.I).strip()
    text = re.sub(r"^(?:Perception\s+Area|Close\s+Area|Ranged\s+Area|Area)\s+", "", text, flags=re.I).strip()
    text = re.sub(r"^(?:Alternate\s+Effect|Alternative)\s*:\s*", "", text, flags=re.I).strip()
    patterns = [
        ("damage", "Damage"),
        ("affliction", "Affliction"),
        ("affliction", "Dazzle"),
        ("weaken", "Weaken"),
        ("enhanced_trait", "Enhanced Trait"),
        ("enhanced_trait", "Enhanced"),
        ("enhanced_trait", "Seize Initiative"),
        ("enhanced_trait", "Advantage"),
        ("healing", "Healing"),
        ("environment", "Environment"),
        ("create", "Create"),
        ("move_object", "Move Object"),
        ("movement", "Movement"),
        ("teleport", "Teleport"),
        ("nullify", "Nullify"),
        ("deflect", "Deflect"),
        ("concealment", "Concealment"),
        ("summon", "Summon"),
        ("transform", "Transform"),
        ("transform", "Morph"),
        ("immunity", "Immunity"),
        ("variable", "Variable"),
        ("protection", "Protection"),
        ("regeneration", "Regeneration"),
        ("illusion", "Illusion"),
        ("mind_reading", "Mind Reading"),
        ("insubstantial", "Insubstantial"),
        ("growth", "Growth"),
        ("shrinking", "Shrinking"),
        ("quickness", "Quickness"),
        ("elongation", "Elongation"),
        ("immortality", "Immortality"),
        ("luck_control", "Luck Control"),
        ("extra_limbs", "Extra Limbs"),
        ("senses", "Senses"),
        ("speed", "Speed"),
        ("flight", "Flight"),
        ("burrowing", "Burrowing"),
        ("swimming", "Swimming"),
        ("leaping", "Leaping"),
        ("communication", "Communication"),
        ("comprehend", "Comprehend"),
        ("remote_sensing", "Remote Sensing"),
        ("feature", "Feature"),
    ]
    for effect_type, label in patterns:
        if re.match(rf"^{re.escape(label)}\b", text, re.I):
            return effect_type, label
    return "custom", "Custom"


def parse_targets_for_effect(effect_type: str, row: dict) -> str:
    if effect_type in {"enhanced_trait", "healing", "regeneration", "protection", "concealment", "senses", "speed", "flight", "burrowing", "swimming", "leaping", "movement", "insubstantial", "growth", "shrinking", "quickness", "elongation", "immortality", "extra_limbs", "comprehend", "communication", "remote_sensing"}:
        return "self"
    if effect_type in {"environment", "create", "illusion", "summon"}:
        return "field"
    if clean(row.get("Teleport")):
        return "target"
    return "target"


def parse_stats_from_segment(segment: str, effect_type: str, row: dict) -> list[str]:
    if effect_type == "enhanced_trait":
        text = re.sub(r"^Linked\s+", "", segment, flags=re.I).strip()
        text = re.sub(r"^Enhanced Trait\b", "", text, flags=re.I).strip()
        text = re.sub(r"^Enhanced\b", "", text, flags=re.I).strip()
        before_rank = re.split(r"\d|\(|\[|;", text)[0]
        stats = [normalize_stat(x) for x in re.split(r"&|/|,|\band\b", before_rank) if clean(x)]
        return [x for x in stats if x]
    if effect_type == "weaken":
        text = re.sub(r"^Linked\s+", "", segment, flags=re.I).strip()
        text = re.sub(r"^Weaken\b", "", text, flags=re.I).strip()
        before_rank = re.split(r"\d|\(|\[|;", text)[0]
        stats = [normalize_stat(x) for x in re.split(r"&|/|,|\band\b", before_rank) if clean(x)]
        return [x for x in stats if x]
    buff = clean(row.get("Buff/Debuff (numérico)"))
    if buff:
        return [normalize_stat(x) for x in re.findall(r"[A-Za-z. ]+(?=\s*[+-])", buff)]
    return []


def parse_effect(row: dict, segment: str, index: int) -> dict:
    effect_type, label = effect_type_from_segment(segment)
    extras, flaws, flags = parse_modifiers(segment)
    status_hint = clean(row.get("Status Secundário"))
    resistance_fallback = "thg" if effect_type == "damage" else (
        "fort" if effect_type in {"affliction", "weaken", "transform"} else (
            "will" if effect_type in {"mind_reading", "nullify"} else (
                "dodge" if effect_type in {"move_object", "teleport"} else ""
            )
        )
    )
    effect = {
        "id": f"e{index}",
        "type": effect_type,
        "label": label,
        "raw": segment,
        "rank": parse_rank(segment, label if label != "Custom" else "", "move"),
        "range": parse_range(row, segment),
        "action": "standard",
        "duration": "instant",
        "resistance": parse_resistance(row, segment, resistance_fallback),
        "target": parse_targets_for_effect(effect_type, row),
        "descriptors": [x for x in [clean(row.get("Tipo")), status_hint] if x],
        "extras": extras,
        "flaws": flaws,
        "linked": flags["linked"] or index > 0,
        "secondaryEffect": flags["secondaryEffect"],
        "fades": flags["fades"],
        "unreliable": flags["unreliable"],
        "reaction": flags["reaction"],
        "selective": flags["selective"],
        "limited": flags["limited"],
        "area": parse_area(row, segment),
    }
    stats = parse_stats_from_segment(segment, effect_type, row)
    if stats:
        effect["traits"] = stats
    if effect_type == "affliction":
        effect["conditions"] = parse_affliction_conditions(segment, status_hint)
    if effect_type in {"create", "move_object", "teleport", "summon", "transform", "variable", "nullify", "deflect", "immunity", "environment", "illusion", "mind_reading", "growth", "shrinking", "luck_control"}:
        effect["requiresChoice"] = effect_type in {"create", "move_object", "teleport", "summon", "transform", "variable", "illusion", "mind_reading", "growth", "shrinking", "luck_control"}
    return effect


def derive_targeting(effects: list[dict], row: dict) -> dict:
    if any(effect.get("area") for effect in effects) or clean(row.get("Área")):
        mode = "area"
    elif effects and all(effect.get("target") == "self" for effect in effects):
        mode = "self"
    elif effects and all(effect.get("target") == "field" for effect in effects):
        mode = "field"
    else:
        mode = "target"
    return {
        "mode": mode,
        "range": "ranged" if is_yes(row.get("Ranged")) else ("perception" if is_yes(row.get("Perception Area")) else "close"),
        "area": parse_area(row, clean(row.get("Build M&M (adaptado)"))),
    }


def build_rule(row: dict, row_number: int, audit_by_name: dict[str, dict]) -> dict:
    name = clean(row.get("Nome")) or f"Move {row_number}"
    build = clean(row.get("Build M&M (adaptado)"))
    if not build:
        build = clean(row.get("Descricao")) or "Damage (Rank = )"
    segments = merge_modifier_only_segments(split_top_level(build))
    effects = [parse_effect(row, seg, idx) for idx, seg in enumerate(segments)]

    if not effects:
        effects = [parse_effect(row, "Damage (Rank = )", 0)]

    audit = audit_by_name.get(slugify(name), {})
    rule_id = f"move:{slugify(name)}:{row_number}"
    extras = sorted({extra for effect in effects for extra in effect.get("extras", [])})
    flaws = sorted({flaw for effect in effects for flaw in effect.get("flaws", [])})
    flags = {
        "linkedEffects": any(effect.get("linked") for effect in effects[1:]),
        "secondaryEffect": any(effect.get("secondaryEffect") for effect in effects),
        "limited": any(effect.get("limited") for effect in effects),
        "fades": any(effect.get("fades") for effect in effects),
        "unreliable": any(effect.get("unreliable") for effect in effects) or norm_text(row.get("Accuracy")) == "unreliable",
        "reaction": any(effect.get("reaction") for effect in effects),
        "selective": any(effect.get("selective") for effect in effects),
    }
    return {
        "schema": "PowerRule",
        "schemaVersion": 1,
        "id": rule_id,
        "sourceRow": row_number,
        "name": name,
        "tm": clean(row.get("TM")),
        "type": clean(row.get("Tipo")),
        "category": clean(row.get("Categoria")),
        "rank": {"source": "move", "value": None},
        "range": derive_targeting(effects, row)["range"],
        "action": "standard",
        "duration": "instant",
        "resistance": normalize_resistance(clean(row.get("Resist Stat"))),
        "descriptors": [x for x in [clean(row.get("Tipo")), clean(row.get("Status Secundário"))] if x],
        "targeting": derive_targeting(effects, row),
        "area": derive_targeting(effects, row)["area"],
        "extras": extras,
        "flaws": flaws,
        "flags": flags,
        "effects": effects,
        "linkedEffects": [effect["id"] for effect in effects if effect.get("linked")],
        "buildText": build,
        "rulesText": clean(row.get("Como funciona (regras/condições)")),
        "description": clean(row.get("Descricao")),
        "audit": {
            "status": clean(audit.get("Veredito")),
            "build": clean(audit.get("Build auditada")),
            "notes": clean(audit.get("Notas da auditoria")),
        },
        "requiresChoices": sorted({
            effect["type"] for effect in effects if effect.get("requiresChoice")
        }),
    }


def read_audit(ws) -> dict[str, dict]:
    headers = [cell.value for cell in ws[4]]
    audit: dict[str, dict] = {}
    for values in ws.iter_rows(min_row=5, values_only=True):
        row = dict(zip(headers, values))
        name = clean(row.get("Nome"))
        if name:
            audit[slugify(name)] = row
    return audit


def build_catalog(source: Path) -> dict:
    workbook = load_workbook(source, data_only=False, read_only=False)
    moves_ws = workbook["Golpes_MM"]
    audit = read_audit(workbook["VALIDAR"]) if "VALIDAR" in workbook.sheetnames else {}
    headers = [cell.value for cell in moves_ws[1]]
    moves = []
    by_name: dict[str, list[str]] = defaultdict(list)
    errors = []
    for row_number, values in enumerate(moves_ws.iter_rows(min_row=2, values_only=True), start=2):
        row = dict(zip(headers, values))
        try:
            rule = build_rule(row, row_number, audit)
            moves.append(rule)
            by_name[slugify(rule["name"])].append(rule["id"])
        except Exception as exc:
            errors.append({"row": row_number, "name": clean(row.get("Nome")), "error": str(exc)})
    return {
        "schema": "MmPowerCatalog",
        "schemaVersion": 1,
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "source": str(source),
        "moveCount": len(moves),
        "errorCount": len(errors),
        "errors": errors,
        "byName": dict(by_name),
        "moves": moves,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", default="../Ga-AlDex/golpes_pokemon_MM_reescritos.xlsx")
    parser.add_argument("--output", default="assets/rules/moves-mm.json")
    args = parser.parse_args()

    source = Path(args.source)
    if not source.exists():
        source = Path("golpes_pokemon_MM_reescritos.xlsx")
    if not source.exists():
        raise FileNotFoundError(args.source)

    catalog = build_catalog(source)
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(catalog, ensure_ascii=True, indent=2), encoding="utf-8")
    print(f"wrote {output} moves={catalog['moveCount']} errors={catalog['errorCount']}")
    return 0 if catalog["errorCount"] == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
